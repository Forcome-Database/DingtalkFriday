"""Statistics regressions using isolated SQLite data and explicit expectations."""

import math
import sys
import types
import unittest
from contextlib import ExitStack
from datetime import date, datetime
from types import SimpleNamespace
from unittest.mock import patch

from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

# Keep tests independent of local credentials and the production database.
test_config = types.ModuleType("app.config")
test_config.settings = SimpleNamespace(database_url="sqlite+aiosqlite:///:memory:")
original_config = sys.modules.get("app.config")
sys.modules["app.config"] = test_config
try:
    from app.database import Base
    from app.models import Department, Employee, LeaveRecord, LeaveType, TripRecord
    from app.services import analytics, leave, trip, trip_analytics
    from app.services.durations import BUSINESS_TIMEZONE, allocate_hours_by_date
    from app.services.export import export_leave_data, export_trip_excel
finally:
    if original_config is None:
        sys.modules.pop("app.config", None)
    else:
        sys.modules["app.config"] = original_config


def instant(value):
    return datetime.fromisoformat(value).replace(tzinfo=BUSINESS_TIMEZONE)


def milliseconds(value):
    return int(instant(value).timestamp() * 1000)


def leave_record(userid, start, end, hours, leave_type="事假", **kwargs):
    return LeaveRecord(
        userid=userid,
        start_time=milliseconds(start),
        end_time=milliseconds(end),
        duration_percent=round(hours * 100),
        duration_unit="percent_hour",
        leave_type=leave_type,
        status="已审批",
        **kwargs,
    )


def trip_record(userid, day, hours, proc, tag="出差"):
    return TripRecord(
        userid=userid,
        work_date=day,
        tag_name=tag,
        begin_time=f"{day} 09:00:00",
        end_time=f"{day} 18:00:00",
        duration_hours=hours,
        proc_inst_id=proc,
        last_synced_at=datetime(2026, 9, 30),
    )


class DurationAllocationTests(unittest.TestCase):
    def test_hour_duration_weights_preserve_cross_month_boundaries(self):
        result = allocate_hours_by_date(instant("2026-06-30 16:00"), instant("2026-07-03 18:00"), 26)
        self.assertEqual(result, {date(2026, 6, 30): 2, date(2026, 7, 1): 8, date(2026, 7, 2): 8, date(2026, 7, 3): 8})

    def test_day_unit_afternoon_half_day(self):
        result = allocate_hours_by_date(instant("2026-09-14 13:30"), instant("2026-09-15 18:00"), 12, calendar_days=True, day_unit=True)
        self.assertEqual(result, {date(2026, 9, 14): 4, date(2026, 9, 15): 8})

    def test_long_day_unit_approval_has_two_half_day_boundaries(self):
        result = allocate_hours_by_date(instant("2026-07-31 13:30"), instant("2026-10-29 13:30"), 720, calendar_days=True, day_unit=True)
        self.assertEqual(len(result), 91)
        self.assertEqual(result[date(2026, 7, 31)], 4)
        self.assertEqual(result[date(2026, 10, 29)], 4)
        self.assertTrue(all(value == 8 for day, value in result.items() if day not in {date(2026, 7, 31), date(2026, 10, 29)}))
        self.assertEqual(math.fsum(result.values()), 720)

    def test_midnight_end_is_exclusive(self):
        result = allocate_hours_by_date(instant("2026-09-30 09:00"), instant("2026-10-01 00:00"), 8)
        self.assertEqual(result, {date(2026, 9, 30): 8})

    def test_same_day_keeps_authoritative_source_hours(self):
        result = allocate_hours_by_date(instant("2026-09-30 09:00"), instant("2026-09-30 18:00"), 6.5)
        self.assertEqual(result, {date(2026, 9, 30): 6.5})

    def test_known_weekend_quota_survives_unmodelled_shift_with_warning(self):
        with self.assertLogs("app.services.durations", level="WARNING") as logs:
            result = allocate_hours_by_date(instant("2026-09-26 09:00"), instant("2026-09-27 18:00"), 16)
        self.assertEqual(result, {date(2026, 9, 26): 8, date(2026, 9, 27): 8})
        self.assertIn("rule conflict", logs.output[0])

    def test_invalid_source_does_not_create_default_hours(self):
        for hours in (-1, float("nan"), float("inf")):
            with self.subTest(hours=hours), self.assertRaises(ValueError):
                allocate_hours_by_date(instant("2026-09-30 09:00"), instant("2026-09-30 18:00"), hours)


class StatisticsTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        async with self.engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        self.patches = ExitStack()
        for module in (analytics, leave, trip, trip_analytics):
            self.patches.enter_context(patch.object(module, "async_session", self.sessions))
        self.patches.enter_context(patch.object(leave, "_is_workday", side_effect=lambda value: value.weekday() < 5))
        self.patches.enter_context(patch.object(leave, "business_today", return_value=date(2026, 9, 30)))
        self.patches.enter_context(patch.object(trip, "business_today", return_value=date(2026, 9, 30)))
        await self.save(
            Department(dept_id=1, name="Company", parent_id=None),
            Department(dept_id=2, name="Sales", parent_id=1),
            Department(dept_id=3, name="Sales", parent_id=1),
            Employee(userid="a", name="Alice", dept_id=2, dept_name="Sales"),
            Employee(userid="b", name="Bob", dept_id=3, dept_name="Sales"),
        )

    async def asyncTearDown(self):
        self.patches.close()
        await self.engine.dispose()

    async def save(self, *rows):
        async with self.sessions() as session:
            session.add_all(rows)
            await session.commit()

    async def test_leave_month_details_and_today_share_26_hour_allocation(self):
        await self.save(leave_record("a", "2026-06-30 16:00", "2026-07-03 18:00", 26))
        table = await leave.get_monthly_summary(2026, unit="hour")
        self.assertEqual(table["list"][0]["months"][5:7], [2, 24])
        self.assertEqual(table["summary"]["total"], 26)
        june = await leave.get_daily_detail("a", 2026, 6)
        july = await leave.get_daily_detail("a", 2026, 7)
        self.assertEqual([row["hours"] for row in june["records"]], [2])
        self.assertEqual([row["hours"] for row in july["records"]], [8, 8, 8])
        self.assertEqual(june["summary"]["totalHours"] + july["summary"]["totalHours"], 26)
        today = await leave.get_today_leave_detail(target_date=date(2026, 6, 30))
        self.assertEqual(today["records"][0]["durationDisplay"], "2小时")
        trend = await analytics.get_monthly_trend(2026)
        self.assertEqual(trend["currentYear"][5:7], [{"month": 6, "days": 0.2}, {"month": 7, "days": 3}])

    async def test_source_day_quota_and_calendar_days_keep_configured_hours(self):
        await self.save(
            LeaveType(leave_code="maternity", leave_name="产假", hours_in_per_day=1200),
            LeaveRecord(userid="a", start_time=milliseconds("2026-09-26 09:00"), end_time=milliseconds("2026-09-28 18:00"), duration_percent=300, duration_unit="percent_day", leave_type="产假", leave_code="maternity", status="已审批"),
        )
        details = await leave.get_daily_detail("a", 2026, 9)
        self.assertEqual([row["hours"] for row in details["records"]], [12, 12, 12])
        self.assertEqual(details["summary"], {"totalDays": 4.5, "totalHours": 36})
        counts = await leave.get_daily_leave_count(2026, 9)
        self.assertEqual([row["count"] for row in counts["days"]][25:28], [1, 1, 1])

    async def test_midnight_end_is_not_counted_in_next_month(self):
        await self.save(leave_record("a", "2026-09-30 09:00", "2026-10-01 00:00", 8))
        october = await leave.get_daily_leave_count(2026, 10)
        self.assertTrue(all(row["count"] == 0 for row in october["days"]))
        details = await leave.get_today_leave_detail(target_date=date(2026, 10, 1))
        self.assertEqual(details, {"date": "2026-10-01", "count": 0, "records": []})

    async def test_overlapping_approvals_remain_additive_but_people_are_distinct(self):
        await self.save(
            leave_record("a", "2026-09-30 09:00", "2026-09-30 18:00", 8),
            leave_record("a", "2026-09-30 10:00", "2026-09-30 17:00", 6),
        )
        summary = await leave.get_monthly_summary(2026, unit="hour")
        self.assertEqual(summary["stats"]["totalCount"], 2)
        self.assertEqual(summary["stats"]["totalDays"], 14)
        counts = await leave.get_daily_leave_count(2026, 9)
        self.assertEqual(counts["todayCount"], 1)
        today = await leave.get_today_leave_detail()
        self.assertEqual(today["count"], 1)
        self.assertEqual(len(today["records"]), 2)

    async def test_empty_types_mean_none_and_pending_records_never_count(self):
        approved = leave_record("a", "2026-09-30 09:00", "2026-09-30 18:00", 8, "年假")
        pending = leave_record("b", "2026-09-30 09:00", "2026-09-30 18:00", 8)
        pending.status = "待复核"
        await self.save(approved, pending)
        self.assertEqual((await leave.get_monthly_summary(2026))["stats"]["totalCount"], 1)
        self.assertEqual((await leave.get_monthly_summary(2026, leave_types=[]))["list"], [])
        self.assertEqual((await leave.get_daily_detail("a", 2026, 9, leave_types=[]))["records"], [])
        self.assertEqual((await leave.get_daily_leave_count(2026, 9, leave_types=[]))["todayCount"], 0)
        self.assertEqual((await leave.get_today_leave_detail(leave_types=[]))["count"], 0)
        self.assertEqual((await analytics.get_leave_type_distribution(2026))["total"], 1)
        self.assertEqual((await analytics.get_employee_ranking(2026))["employees"][0]["name"], "Alice")

    async def test_leave_detail_inherits_type_filter(self):
        await self.save(
            leave_record("a", "2026-09-29 09:00", "2026-09-29 18:00", 8, "年假"),
            leave_record("a", "2026-09-30 09:00", "2026-09-30 18:00", 8, "事假"),
        )
        detail = await leave.get_daily_detail("a", 2026, 9, leave_types=["年假"])
        self.assertEqual([row["leaveType"] for row in detail["records"]], ["年假"])
        self.assertEqual(detail["summary"]["totalHours"], 8)

    async def test_leave_totals_aggregate_before_rounding(self):
        await self.save(
            leave_record("a", "2026-09-29 09:00", "2026-09-29 09:30", 0.5),
            leave_record("b", "2026-09-29 09:00", "2026-09-29 09:30", 0.5),
        )
        result = await leave.get_monthly_summary(2026)
        self.assertEqual(result["stats"]["totalDays"], 0.1)
        self.assertEqual(result["summary"]["total"], 0.1)
        self.assertEqual(result["summary"]["months"][8], 0.1)

    async def test_trip_short_hours_and_shared_process_keep_person_counts(self):
        await self.save(
            trip_record("a", "2026-09-29", 0.5, "same-process"),
            trip_record("a", "2026-09-30", 0.5, "same-process"),
            trip_record("b", "2026-09-29", 0.5, "same-process"),
        )
        result = await trip.get_trip_monthly_summary(2026)
        self.assertEqual(result["stats"]["totalCount"], 2)
        self.assertEqual(result["stats"]["totalDays"], 0.2)
        self.assertEqual(result["summary"]["totalDays"], 0.2)
        self.assertEqual(result["summary"]["months"]["9"], 0.2)
        self.assertEqual(result["list"][0]["months"]["9"], 0.1)

    async def test_december_31_detail_and_trip_type_filter(self):
        await self.save(trip_record("a", "2026-12-31", 8, "trip"), trip_record("a", "2026-12-30", 4, "outing", tag="外出"))
        detail = await trip.get_trip_daily_detail("a", 2026, 12, trip_type="出差")
        self.assertEqual([row["date"] for row in detail["records"]], ["2026-12-31"])
        self.assertEqual(detail["summary"], {"totalDays": 1, "totalHours": 8})
        table = await trip.get_trip_monthly_summary(2026, trip_type="出差")
        self.assertEqual(table["list"][0]["months"]["12"], 1)

    async def test_zero_hour_trip_does_not_count_as_person(self):
        await self.save(trip_record("a", "2026-09-30", 0, "empty"), trip_record("b", "2026-09-30", 0, "outing", tag="外出"))
        summary = await trip.get_trip_monthly_summary(2026)
        self.assertEqual(summary["stats"]["todayTripCount"], 0)
        self.assertEqual(summary["stats"]["todayOutingCount"], 0)
        self.assertEqual(summary["stats"]["totalCount"], 0)
        self.assertEqual(summary["list"], [])
        self.assertEqual((await trip.get_trip_daily_detail("a", 2026, 9))["records"], [])
        self.assertEqual((await trip.get_trip_today())["list"], [])
        self.assertEqual((await trip.get_trip_daily_count(2026, 9))["days"], {})
        self.assertEqual((await trip_analytics.get_trip_employee_ranking(2026))["employees"], [])
        self.assertTrue(all(row["count"] == 0 for row in (await trip_analytics.get_trip_weekday_distribution(2026))["weekdays"]))
        departments = (await trip_analytics.get_trip_department_comparison(2026))["departments"]
        self.assertEqual(sum(row["headcount"] for row in departments), 2)

    async def test_zero_trip_records_do_not_increment_valid_employee_counts(self):
        await self.save(
            trip_record("a", "2026-09-30", 8, "valid"),
            trip_record("a", "2026-09-29", 0, "empty-trip"),
            trip_record("a", "2026-09-30", 0, "empty-outing", tag="外出"),
        )
        row = (await trip.get_trip_monthly_summary(2026))["list"][0]
        self.assertEqual(row["tripCount"], 1)
        self.assertEqual(row["outingCount"], 0)
        self.assertEqual(row["totalDays"], 1)

    async def test_today_total_deduplicates_people_across_trip_types(self):
        await self.save(
            trip_record("a", "2026-09-30", 4, "trip"),
            trip_record("a", "2026-09-30", 4, "outing", tag="外出"),
        )
        stats = (await trip.get_trip_monthly_summary(2026))["stats"]
        self.assertEqual(stats["todayTripCount"], 1)
        self.assertEqual(stats["todayOutingCount"], 1)
        self.assertEqual(stats["todayTotalCount"], 1)
        filtered = (await trip.get_trip_monthly_summary(2026, trip_type="出差"))["stats"]
        self.assertEqual(filtered["todayOutingCount"], 0)
        self.assertEqual(filtered["todayTotalCount"], 1)

    async def test_departments_are_distinct_by_id_and_mean_uses_requested_metric(self):
        await self.save(
            Employee(userid="c", name="Carol", dept_id=2, dept_name="Sales"),
            leave_record("a", "2026-09-30 09:00", "2026-09-30 18:00", 8),
            leave_record("b", "2026-09-30 09:00", "2026-09-30 18:00", 24),
            trip_record("a", "2026-09-30", 8, "one"),
            trip_record("b", "2026-09-30", 24, "two"),
        )
        for operation in (analytics.get_department_comparison, trip_analytics.get_trip_department_comparison):
            with self.subTest(operation=operation.__name__):
                total = await operation(2026, metric="total")
                average = await operation(2026, metric="avg")
                self.assertEqual(len(total["departments"]), 2)
                self.assertEqual({row["name"] for row in total["departments"]}, {"Company / Sales [2]", "Company / Sales [3]"})
                self.assertEqual(total["average"], 2)
                self.assertEqual(average["average"], 1.8)
                self.assertEqual({row["headcount"] for row in total["departments"]}, {1, 2})

    async def test_cross_year_preserves_original_hours(self):
        await self.save(leave_record("a", "2025-12-31 16:00", "2026-01-02 18:00", 18))
        old = await leave.get_monthly_summary(2025, unit="hour")
        new = await leave.get_monthly_summary(2026, unit="hour")
        self.assertEqual(old["summary"]["total"], 2)
        self.assertEqual(new["summary"]["total"], 16)
        self.assertEqual((await leave.get_daily_detail("a", 2026, 1))["summary"]["totalHours"], 16)

    async def test_excel_keeps_filtered_months_and_totals(self):
        from openpyxl import load_workbook

        await self.save(
            leave_record("a", "2026-06-30 16:00", "2026-07-03 18:00", 26, "年假"),
            trip_record("a", "2026-12-31", 8, "trip"),
            trip_record("a", "2026-12-30", 4, "outing", tag="外出"),
        )
        leave_sheet = load_workbook(await export_leave_data(2026, leave_types=["年假"], unit="hour")).active
        self.assertEqual([leave_sheet.cell(2, column).value for column in (8, 9, 15)], [2, 24, 26])
        trip_sheet = load_workbook(await export_trip_excel(2026, trip_type="出差")).active
        self.assertEqual([trip_sheet.cell(2, column).value for column in (3, 4, 16, 17)], [1, 0, 1, 1])


if __name__ == "__main__":
    unittest.main()
